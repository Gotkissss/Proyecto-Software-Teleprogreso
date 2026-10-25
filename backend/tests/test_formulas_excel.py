"""HU-S10-04, tarea 6: un texto que empieza por = + - @ aparece como texto en todos los XLSX."""
from datetime import date
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.services.exportacion_reportes import (
    _crear_libro,
    generar_excel_tabla,
    texto_seguro_excel,
)

PELIGROSOS = [
    '=HYPERLINK("http://malo.example","clic")',
    "=1+1",
    "+cmd|' /C calc'!A0",
    "-2+3",
    "@SUM(1,1)",
    "\t=1+1",
    "\r=1+1",
]


@pytest.mark.parametrize("texto", PELIGROSOS)
def test_texto_peligroso_recibe_apostrofo(texto):
    assert texto_seguro_excel(texto) == "'" + texto


@pytest.mark.parametrize("valor", ["Instalar router", "", "a=b", " =1", 5, -5, 0.5, None, date(2026, 10, 24)])
def test_lo_demas_pasa_sin_tocar(valor):
    assert texto_seguro_excel(valor) == valor


def _celdas(contenido: bytes):
    hoja = load_workbook(BytesIO(contenido)).active
    return [c for fila in hoja.iter_rows() for c in fila if c.value is not None]


def _sin_formulas(celdas):
    return [c.coordinate for c in celdas if c.data_type == "f" or str(c.value).startswith(("=", "+", "@"))]


def test_generar_excel_tabla_no_deja_formulas_en_ninguna_celda():
    contenido = generar_excel_tabla(
        titulo=PELIGROSOS[0],
        subtitulo=PELIGROSOS[4],
        nombre_hoja="Bitacora",
        encabezados=("accion", PELIGROSOS[1]),
        filas=[(p, -7) for p in PELIGROSOS],
    )
    celdas = _celdas(contenido)

    assert _sin_formulas(celdas) == []
    # El número negativo sigue siendo número, no texto.
    assert all(c.data_type == "n" for c in celdas if c.value == -7)


def test_crear_libro_no_deja_formulas_en_titulo_filas_ni_totales():
    reporte = SimpleNamespace(fecha_inicio=date(2026, 10, 1), fecha_fin=date(2026, 10, 24))
    contenido = _crear_libro(
        titulo=PELIGROSOS[0],
        nombre_hoja="Reporte",
        reporte=reporte,
        encabezados=("empleado", "total"),
        filas=[(p, 3) for p in PELIGROSOS],
        fila_total=("=SUM(B5:B11)", 21),
    )

    assert _sin_formulas(_celdas(contenido)) == []
